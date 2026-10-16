# -*- coding: utf-8 -*-
"""
Mintaadat-generátor a Termék kereső bemutatóhoz.

Létrehozza a 'minta_adatok' mappát kitalált termékekkel, ügyfelekkel és fájlokkal.
Az adatok szándékosan "valósághűen rendetlenek": ugyanazokat a csapdákat tartalmazzák,
amelyekkel egy valódi üzemi fájlrendszerben találkozni lehet (lásd a README csapda-táblázatát).

Futtatás:  python mintaadat_general.py      (újrafuttatható, mindig újra felépíti)
"""
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ALAP = Path(__file__).parent / "minta_adatok"

# ---- Termékek (Excel sorok) -------------------------------------------------
# termék, gép, programkód, szerszám, magasság (mm), anyag, megjegyzés, változat
TERMEKEK = [
    ("1042",       "1", "P10421",               "SZ-12", 35, "acél 2 mm",    "",                                   ""),
    ("1042",       "3", "10423",                "SZ-12", 35, "acél 2 mm",    "3-as gépen előtag nélküli kód",      ""),
    ("10421",      "2", "P104211",              "SZ-07", 40, "acél 1,5 mm",  "hasonló szám, MÁS termék",           ""),
    ("2210 X",     "1", "P22101",               "SZ-03", 28, "alumínium",    "toldalék a cellában",                ""),
    ("3300",       "2", "P33001",               "SZ-21", 52, "rozsdamentes", "",                                   "3301 bevonatos"),
    ("450 120 00", "4", "45012",                "SZ-09", 31, "acél 3 mm",    "szóköz a számban, kötőjel a fájlban", ""),
    ("5005",       "3", "50051; 50052; 50055",  "SZ-14", 44, "acél 2 mm",    "több program, egy hiányzik",         ""),
    ("5100",       "4", "51001 ?",              "SZ-02", 25, "alumínium",    "bizonytalan kód",                    ""),
    ("6020",       "1", "P60201",               "SZ-11", 38, "acél 2 mm",    "program rossz gépnél",               ""),
    ("7000",       "2", "",                     "SZ-05", 30, "",             "még nincs program",                  ""),
]
# + rendes, "unalmas" termékek, hogy legyen tömeg
for i, n in enumerate(range(8001, 8021)):
    gep = str(i % 4 + 1)
    kod = f"P{n}1" if gep in "12" else f"{n}1"
    TERMEKEK.append((str(n), gep, kod, f"SZ-{10 + i % 9:02d}", 30 + i % 7 * 3,
                     ["acél 2 mm", "alumínium", "rozsdamentes"][i % 3], "", ""))

UGYFELEK = ["Alfa Kft", "Béta Zrt", "Gamma Bt", "Delta GmbH"]

# ---- Fájlok -----------------------------------------------------------------
# (ügyfél mappa, fájlnév) - termékadatlapok
ADATLAPOK = [
    ("Alfa Kft", "1042_A_HU-EN.pdf"),
    ("Alfa Kft/ARCHÍV", "1042_regi.pdf"),             # archív -> nem jelenhet meg
    ("Béta Zrt", "10421_B.pdf"),                        # nem jelenhet meg 1042-nél
    ("Alfa Kft", "2210_C_HU.pdf"),
    ("Gamma Bt", "3300_A_HU-EN.pdf"),
    ("Delta GmbH", "450 120-00_A_HU-EN.pdf"),
    ("Delta GmbH/HU", "450 120-00_A_HU.pdf"),           # almappa: ugyanaz az ügyfél
    ("Béta Zrt", "5005_B.pdf"),
    ("Gamma Bt", "5100_A.pdf"),
    ("Alfa Kft", "6020_A.pdf"),
    ("Béta Zrt", "7000_C.pdf"),
]
ADATLAPOK += [(UGYFELEK[i % 4], f"{n}_A.pdf") for i, n in enumerate(range(8001, 8021))]

# 3D modellek - szándékosan nem mindegyik termékhez van
MODELLEK = ["1042_A.stp", "10421.step", "3300_bevonat_nelkul.stp", "450 120-00.stp",
            "5005.STEP", "8001.stp", "8002.stp", "8005.stp", "8010.stp", "8011.stp"]

# Gépi programok: gép mappa -> fájlok
PROGRAMOK = {
    "programok/Gep-1": ["P10421.prg", "P22101.prg", "P80011.prg", "P80051.prg", "P80091.prg",
                        "P80131.prg", "P80171.prg"],
    "programok/Gep-2": ["P104211.prg", "P33001.prg", "P60201.prg",          # P60201: rossz gépnél!
                        "P80021.prg", "P80061.prg", "P80101.prg", "P80141.prg", "P80181.prg"],
    "programok/Gep-3": ["P10423.prg", "P50051.prg", "P50052.prg",           # P50055 hiányzik
                        "P80031.prg", "P80071.prg", "P80111.prg", "P80151.prg", "P80191.prg"],
    "programok/Gep-4": ["P45012.prg", "P80041.prg", "P80081.prg", "P80121.prg", "P80161.prg"],
    "programok/Gep-4/ARCHÍV": ["P45012.prg"],                                # archív -> kimarad
    "programok_mentes/Gep-4": ["P51001.prg", "P80201.prg"],                  # gép mentése
}


def pdf_fajl(ut, sorok):
    """Egyoldalas, szabványos PDF külső könyvtár nélkül."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    szoveg = "BT /F1 18 Tf 72 760 Td (" + esc(sorok[0]) + ") Tj /F1 11 Tf"
    for s in sorok[1:]:
        szoveg += " 0 -22 Td (" + esc(s) + ") Tj"
    szoveg += " ET"
    stream = szoveg.encode("latin-1", "replace")
    objektumok = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    ki, poz = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objektumok, 1):
        poz.append(len(ki))
        ki += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(ki)
    ki += f"xref\n0 {len(objektumok) + 1}\n0000000000 65535 f \n".encode()
    for p in poz:
        ki += f"{p:010d} 00000 n \n".encode()
    ki += f"trailer\n<< /Size {len(objektumok) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    ut.write_bytes(bytes(ki))


def step_fajl(ut):
    ut.write_text(
        "ISO-10303-21;\nHEADER;\n"
        f"FILE_DESCRIPTION(('Mintamodell - {ut.stem}'),'2;1');\n"
        f"FILE_NAME('{ut.name}','2026-01-01T00:00:00',('demo'),(''),'','','');\n"
        "FILE_SCHEMA(('AUTOMOTIVE_DESIGN'));\nENDSEC;\nDATA;\nENDSEC;\nEND-ISO-10303-21;\n",
        encoding="utf-8")


def program_fajl(ut, gep):
    ut.write_text(f"; Mintaprogram {ut.stem}\n; Gép: {gep}\n; (kitalált tartalom)\n"
                  "G90\nG0 X0 Y0\nM30\n", encoding="utf-8")


def main():
    if ALAP.exists():
        shutil.rmtree(ALAP)
    ALAP.mkdir()

    # Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Termékek"
    fejlec = ["termék", "gép", "programkód", "szerszám", "beállítási magasság (mm)",
              "anyag", "megjegyzés", "változat"]
    ws.append(fejlec)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
    for t in TERMEKEK:
        ws.append([int(x) if isinstance(x, str) and x.isdigit() and i in (0, 1) else x
                   for i, x in enumerate(t)])
    for i, w in enumerate([14, 6, 22, 10, 24, 14, 36, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ALAP / "termekek.xlsx")

    for mappa, nev in ADATLAPOK:
        d = ALAP / "adatlapok" / mappa
        d.mkdir(parents=True, exist_ok=True)
        ugyfel = mappa.split("/")[0]
        pdf_fajl(d / nev, [f"Termekadatlap - {Path(nev).stem}",
                           f"Ugyfel: {ugyfel}", "Kitalalt mintadokumentum (demo)"])

    d = ALAP / "3d_modellek"
    d.mkdir(parents=True)
    for nev in MODELLEK:
        step_fajl(d / nev)

    for mappa, fajlok in PROGRAMOK.items():
        d = ALAP / mappa
        d.mkdir(parents=True, exist_ok=True)
        for nev in fajlok:
            program_fajl(d / nev, mappa.split("/")[1])

    db = sum(1 for _ in ALAP.rglob("*") if _.is_file())
    print(f"Kész: {ALAP}  ({len(TERMEKEK)} Excel-sor, {db} fájl)")


if __name__ == "__main__":
    main()
