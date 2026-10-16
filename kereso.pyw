# -*- coding: utf-8 -*-
"""
Termék kereső - egy termékszámra egy ablakban mutatja a termék beállításait
és a hozzá tartozó fájlokat (adatlap, 3D modell, gépi program) több forrásból.

Beállítás:  config.json (mappák, oszlopnevek, gépek)
Telepítés:  pip install openpyxl
Futtatás:   kereso.pyw (Windows: dupla kattintás)  |  python kereso.pyw
"""
import json
import os
import re
import subprocess
import sys
import threading
import warnings
import tkinter as tk
from tkinter import ttk, messagebox
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

APP_MAPPA = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).parent


# ---- Beállítások -------------------------------------------------------------
def config_betolt():
    """config.json beolvasása; a relatív útvonalakat az alkalmazás mappájához igazítja."""
    with open(APP_MAPPA / "config.json", encoding="utf-8") as f:
        c = json.load(f)

    def ut(p):
        p = Path(p)
        return p if p.is_absolute() else APP_MAPPA / p

    c["excel"] = ut(c["excel"])
    for fo in c["forrasok"]:
        fo["mappa"] = ut(fo["mappa"])
        fo["kiterjesztesek"] = tuple(k.lower() for k in fo["kiterjesztesek"])
    pr = c.get("programok")
    if pr:
        pr["kiterjesztesek"] = tuple(k.lower() for k in pr["kiterjesztesek"])
        for g in pr["gepek"].values():
            g["mappak"] = [ut(m) for m in g["mappak"]]
    c["kihagyott_mappak"] = [k.lower() for k in c.get("kihagyott_mappak", [])]
    return c


try:
    CFG = config_betolt()
except Exception as _e:
    CFG = None
    _CFG_HIBA = _e


# ---- Szövegkezelés és párosítás ---------------------------------------------
def cella_szoveg(v):
    """Cella érték szöveggé. 12345.0 -> '12345' (Excel a számokat float-ként adja)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def norm(szoveg):
    """Összehasonlításhoz: kisbetű, minden szóköz törölve. '450 120' == '450120'"""
    return "".join(szoveg.split()).lower()


def reszek(szoveg):
    """Részekre bontás szóköz, _, -, pont mentén.
    '450 120-00_A_HU-EN' -> ['450', '120', '00', 'a', 'hu', 'en']"""
    return re.findall(r"[^\W_]+", szoveg.lower())


def termekszam_mag(termek):
    """A cellából a termékszám: első rész + az utána következő csak számjegyes részek.
    '2210 X' -> ['2210'];  '450 120 00' -> ['450', '120', '00']"""
    r = reszek(termek)
    if not r:
        return []
    mag = [r[0]]
    for x in r[1:]:
        if not x.isdigit():
            break
        mag.append(x)
    return mag


def illik(termek, fajl_reszek):
    """A fájlnév a termékszámmal kezdődik, és utána nem folytatódik betűvel/számmal.
    '1042' -> 1042_A_HU-EN.pdf igen;  '1042' -> 10421_B.pdf NEM (más termék)"""
    mag = termekszam_mag(termek)
    return bool(mag) and list(fajl_reszek[:len(mag)]) == mag


def kihagyando(mappanev):
    return any(k in mappanev.lower() for k in CFG["kihagyott_mappak"])


def fajl_index(mappa, kiterj):
    """[(fájlnév részei, teljes út)], olvashatatlan mappák - archív mappák nélkül."""
    if not mappa.is_dir():
        raise FileNotFoundError(mappa)
    lista, hibak = [], []
    for gyoker, mappak, fajlok in os.walk(mappa, onerror=hibak.append):
        mappak[:] = [m for m in mappak if not kihagyando(m)]
        for f in fajlok:
            alap, kit = os.path.splitext(f)
            if kit.lower() in kiterj:
                lista.append((reszek(alap), Path(gyoker) / f))
    return lista, hibak


def program_index():
    """{programnév kisbetűvel: [(gép kulcs, gép neve, teljes út)]}, hibák száma, fájlok száma"""
    pr = CFG.get("programok")
    idx, hibak, db = {}, [], 0
    if not pr:
        return idx, 0, 0
    for kulcs, gep in pr["gepek"].items():
        for m in gep["mappak"]:
            if not m.is_dir():
                continue
            for gy, almappak, fajlok in os.walk(m, onerror=hibak.append):
                almappak[:] = [a for a in almappak if not kihagyando(a)]
                for f in fajlok:
                    alap, kit = os.path.splitext(f)
                    if kit.lower() in pr["kiterjesztesek"]:
                        idx.setdefault(alap.lower(), []).append((kulcs, gep["nev"], Path(gy) / f))
                        db += 1
    return idx, len(hibak), db


def programkodok(cella):
    """Programkód cellából a keresendő nevek kódonként.
    'P10421' -> [['p10421']];  '10423' -> [['10423', 'p10423']] (előtag pótlása);
    '50051; 50052' -> két kód;  '51001 ?' -> bizonytalan
    Vissza: ([(eredeti kód, [jelöltek])], bizonytalan-e)"""
    elotag = CFG["programok"].get("elotag", "").lower()
    bizonytalan = "?" in cella
    kodok = []
    for tok in re.split(r"[;,/\s]+", cella):
        t = tok.strip("?").strip().lower()
        if not t or not any(c.isdigit() for c in t):
            continue
        jeloltek = [t, elotag + t] if (elotag and t[0].isdigit()) else [t]
        kodok.append((tok.strip("?").strip(), jeloltek))
    return kodok, bizonytalan


def rovid_ut(p):
    """Az alkalmazás mappáján belüli fájl relatív úttal, egyébként teljes úttal."""
    try:
        return str(p.relative_to(APP_MAPPA))
    except ValueError:
        return str(p)


def fajl_megnyit(p):
    """Megnyitás a rendszerben hozzárendelt programmal (Windows / macOS / Linux)."""
    try:
        if sys.platform.startswith("win"):
            try:
                os.startfile(p)
            except OSError:
                subprocess.Popen(["notepad.exe", str(p)])   # nincs társított program
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(p)])
        else:
            subprocess.Popen(["xdg-open", str(p)])
    except Exception as e:
        messagebox.showerror("Nem nyitható meg", f"{p}\n\n{e}")


# ---- Excel ------------------------------------------------------------------
def betolt():
    wb = openpyxl.load_workbook(CFG["excel"], read_only=True, data_only=True)
    lap = CFG["munkalap"]
    if lap not in wb.sheetnames:
        nevek = wb.sheetnames
        wb.close()
        raise ValueError(f"Nincs '{lap}' nevű munkalap. Létező lapok: {nevek}")
    sorok = list(wb[lap].iter_rows(values_only=True))
    wb.close()
    if not sorok:
        raise ValueError("A munkalap üres.")

    # Fejléc sor: az első 30 sorból az, amelyikben szerepel valamelyik kulcsszó
    fejlec_idx = kulcs = None
    for szo in CFG["kereso_fejlecek"]:
        for i, s in enumerate(sorok[:30]):
            j = next((j for j, c in enumerate(s) if szo.lower() in cella_szoveg(c).lower()), None)
            if j is not None:
                fejlec_idx, kulcs = i, j
                break
        if fejlec_idx is not None:
            break
    if fejlec_idx is None:
        fejlec_idx = next((i for i, s in enumerate(sorok) if any(c is not None for c in s)), 0)
        kulcs = 0

    fejlec = [cella_szoveg(c) or f"Oszlop {i + 1}" for i, c in enumerate(sorok[fejlec_idx])]
    adat = []
    for excel_sor, s in enumerate(sorok[fejlec_idx + 1:], start=fejlec_idx + 2):
        if not any(c is not None for c in s):
            continue
        ertekek = [cella_szoveg(c) for c in s]
        ertekek += [""] * (len(fejlec) - len(ertekek))
        adat.append((excel_sor, ertekek))
    return fejlec, adat, kulcs


# ---- Felület ------------------------------------------------------------------
class App:
    def __init__(self, root):
        self.root = root
        root.title("Termék kereső")

        felso = ttk.Frame(root, padding=8)
        felso.pack(fill="x")
        ttk.Label(felso, text="Termék:").pack(side="left")
        self.mezo = ttk.Entry(felso, font=("Segoe UI", 12))
        self.mezo.pack(side="left", fill="x", expand=True, padx=6)
        self.mezo.bind("<Return>", lambda e: self.keres())
        ttk.Button(felso, text="Keresés", command=self.keres).pack(side="left")
        ttk.Button(felso, text="Újratöltés", command=self.ujratolt).pack(side="left", padx=(6, 0))

        oszlop_sor = ttk.Frame(root, padding=(8, 0, 8, 6))
        oszlop_sor.pack(fill="x")
        ttk.Label(oszlop_sor, text="Keresés oszlopa:").pack(side="left")
        self.oszlop_valaszto = ttk.Combobox(oszlop_sor, state="readonly", font=("Segoe UI", 11))
        self.oszlop_valaszto.pack(side="left", fill="x", expand=True, padx=6)
        self.oszlop_valaszto.bind("<<ComboboxSelected>>", self.oszlop_valtas)

        self.lista = tk.Listbox(root, height=5, font=("Consolas", 11))
        self.lista.pack(fill="x", padx=8)
        self.lista.bind("<<ListboxSelect>>", self.lista_valaszt)

        # Alsó részek előbb (az eredmény mező kapja a maradék helyet)
        self.allapot = ttk.Label(root, padding=(8, 0, 8, 6))
        self.allapot.pack(side="bottom", fill="x")

        pr = CFG.get("programok")
        self.prog_lista = None
        self.prog_idx, self.prog_talalat, self.prog_allapot = {}, [], ""
        if pr:
            k = ttk.LabelFrame(root, text=f"{pr['nev']} (dupla kattintás = megnyitás)", padding=6)
            k.pack(side="bottom", fill="x", padx=8, pady=(0, 6))
            self.prog_lista = tk.Listbox(k, height=5, font=("Segoe UI", 11))
            self.prog_lista.pack(side="left", fill="x", expand=True)
            self.prog_lista.bind("<Double-Button-1>", lambda e: self.prog_megnyit())
            ttk.Button(k, text="Megnyitás", command=self.prog_megnyit).pack(side="left", padx=(6, 0))

        self.forrasok = []
        for fo in reversed(CFG["forrasok"]):
            keret = ttk.LabelFrame(root, text=f"{fo['nev']} (dupla kattintás = megnyitás)", padding=6)
            keret.pack(side="bottom", fill="x", padx=8, pady=(0, 6))
            lb = tk.Listbox(keret, height=5, font=("Segoe UI", 11))
            lb.pack(side="left", fill="x", expand=True)
            a = {"def": fo, "lista": lb, "index": [], "talalat": [], "allapot": ""}
            lb.bind("<Double-Button-1>", lambda e, a=a: self.megnyit(a))
            ttk.Button(keret, text="Megnyitás", command=lambda a=a: self.megnyit(a)).pack(side="left", padx=(6, 0))
            self.forrasok.insert(0, a)

        self.kimenet = tk.Text(root, font=("Consolas", 12), wrap="word", width=72, height=12)
        self.kimenet.pack(fill="both", expand=True, padx=8, pady=8)
        self.kimenet.tag_configure("cim", font=("Consolas", 12, "bold"))
        self.kimenet.tag_configure("kiemel", font=("Consolas", 12, "bold"), foreground="#1a5fb4")

        self.fejlec, self.adat, self.kulcs = [], [], 0
        self.talalatok, self.valt_jel, self.aktualis = [], {}, None
        self.index_kesz = False
        self.ujratolt()
        self.mezo.focus()

    def oszlop(self, szo):
        """Az első oszlop indexe, amelynek fejlécében szerepel a szó (vagy None)."""
        szo = szo.lower()
        return next((i for i, f in enumerate(self.fejlec) if szo in f.lower()), None)

    # ---- Betöltés ----
    def ujratolt(self):
        try:
            self.fejlec, self.adat, self.kulcs = betolt()
            self.oszlop_valaszto["values"] = [f"{get_column_letter(i + 1)}: {f}"
                                              for i, f in enumerate(self.fejlec)]
            self.oszlop_valaszto.current(self.kulcs)
            valodi = sum(1 for i, f in enumerate(self.fejlec) if f != f"Oszlop {i + 1}")
            self.kimenet.config(height=valodi + 5)
        except Exception as e:
            self.fejlec, self.adat, self.kulcs = [], [], 0
            self.oszlop_valaszto["values"] = []
            messagebox.showerror("Hiba az Excel betöltésénél", f"{CFG['excel']}\n\n{e}")

        self.index_kesz = False
        for f in self.forrasok:
            f["allapot"] = f"{f['def']['nev']}: beolvasás..."
        self.prog_allapot = "programok: beolvasás..." if self.prog_lista else ""
        self.allapot_frissit()
        threading.Thread(target=self._index_szal, daemon=True).start()
        self.root.after(200, self._index_figyel)

    def _index_szal(self):
        for f in self.forrasok:
            d = f["def"]
            try:
                f["index"], hibak = fajl_index(d["mappa"], d["kiterjesztesek"])
                f["allapot"] = f"{d['nev']}: {len(f['index'])}"
                if hibak:
                    f["allapot"] += f" (FIGYELEM: {len(hibak)} mappa nem olvasható)"
            except Exception:
                f["index"] = []
                f["allapot"] = f"{d['nev']}: mappa nem elérhető"
        if self.prog_lista:
            try:
                self.prog_idx, hiba, db = program_index()
                self.prog_allapot = f"{CFG['programok']['nev']}: {db}" + \
                    (f" (FIGYELEM: {hiba} mappa nem olvasható)" if hiba else "")
            except Exception:
                self.prog_idx, self.prog_allapot = {}, "programok: hiba"
        self.index_kesz = True

    def _index_figyel(self):
        if not self.index_kesz:
            self.root.after(200, self._index_figyel)
            return
        self.allapot_frissit()
        if self.aktualis:
            self.fajlok_keres(self.aktualis[1][self.kulcs])
            self.prog_keres(self.aktualis[1])

    def allapot_frissit(self):
        reszlet = " | ".join([f["allapot"] for f in self.forrasok] + ([self.prog_allapot] if self.prog_allapot else []))
        self.allapot.config(text=f"Betöltve: {len(self.adat)} sor | {reszlet}")

    def oszlop_valtas(self, _):
        self.kulcs = self.oszlop_valaszto.current()
        self.keres()

    # ---- Keresés ----
    def keres(self):
        q = norm(self.mezo.get())
        self.lista.delete(0, "end")
        self.kimenet.delete("1.0", "end")
        for f in self.forrasok:
            f["lista"].delete(0, "end")
            f["talalat"] = []
        if self.prog_lista:
            self.prog_lista.delete(0, "end")
        self.prog_talalat, self.aktualis = [], None
        if not q:
            return

        iv = self.oszlop(CFG.get("valtozat_fejlec", "változat"))

        def valtozatok(e):
            if iv is None or not e[iv]:
                return []
            return [(norm(t), t.strip()) for t in re.split(r"[;,/\n]+", e[iv]) if t.strip()]

        # Először pontos egyezés (termék vagy változat), ha nincs, részleges
        for pontos_kell in (True, False):
            self.talalatok, self.valt_jel = [], {}
            for s in self.adat:
                k = norm(s[1][self.kulcs])
                if (k == q) if pontos_kell else (q in k):
                    self.talalatok.append(s)
                    continue
                for vn, veredeti in valtozatok(s[1]):
                    if (vn == q) if pontos_kell else (q in vn):
                        self.talalatok.append(s)
                        self.valt_jel[s[0]] = veredeti
                        break
            if self.talalatok:
                break

        if not self.talalatok:
            self.kimenet.insert("end", "Nincs találat.")
        elif len(self.talalatok) == 1:
            self.mutat(self.talalatok[0])
        else:
            gs = self.oszlop(CFG["programok"]["gep_fejlec"]) if CFG.get("programok") else None
            szel = max(len(e[self.kulcs]) for _, e in self.talalatok)
            for excel_sor, e in self.talalatok:
                info = f"Gép: {e[gs] or '-'}" if gs is not None else f"Excel sor: {excel_sor}"
                if excel_sor in self.valt_jel:
                    info += f"   (változat: {self.valt_jel[excel_sor]})"
                self.lista.insert("end", f"{e[self.kulcs]:<{szel}}   |   {info}")
            self.kimenet.insert("end", f"{len(self.talalatok)} találat - válassz a listából.")

    def lista_valaszt(self, _):
        sel = self.lista.curselection()
        if sel:
            self.mutat(self.talalatok[sel[0]])

    def mutat(self, talalat):
        self.aktualis = talalat
        excel_sor, e = talalat
        self.kimenet.delete("1.0", "end")
        self.kimenet.insert("end", f"Excel sor: {excel_sor}")
        if excel_sor in self.valt_jel:
            self.kimenet.insert("end", f"   |   Változat: {self.valt_jel[excel_sor]}  →  alap: {e[self.kulcs]}", "kiemel")
        self.kimenet.insert("end", "\n\n")
        sorok = [(f, v) for i, (f, v) in enumerate(zip(self.fejlec, e)) if v or f != f"Oszlop {i + 1}"]
        szel = max((len(f) for f, _ in sorok), default=0)
        for f, v in sorok:
            self.kimenet.insert("end", f"{f:<{szel}} : ", "cim")
            self.kimenet.insert("end", f"{v or '-'}\n")
        self.fajlok_keres(e[self.kulcs])
        self.prog_keres(e)

    # ---- Fájlforrások ----
    def fajlok_keres(self, termek):
        csoportok, van_csoport_forras = [], False
        for f in self.forrasok:
            d, lb = f["def"], f["lista"]
            lb.delete(0, "end")
            f["talalat"] = []
            if not self.index_kesz:
                lb.insert("end", "Beolvasás folyamatban...")
                continue
            f["talalat"] = sorted((p for nev, p in f["index"] if illik(termek, nev)),
                                  key=lambda p: str(p).lower())
            if not f["talalat"]:
                lb.insert("end", "Nincs találat.")
            for p in f["talalat"]:
                try:
                    rel = p.relative_to(d["mappa"])
                except ValueError:
                    rel = p
                lb.insert("end", str(rel))
                # csoport (pl. ügyfél) = első szintű mappa neve
                if d.get("csoport") and len(rel.parts) > 1 and rel.parts[0] not in csoportok:
                    csoportok.append(rel.parts[0])
            if f["talalat"]:
                lb.selection_set(0)
            van_csoport_forras = van_csoport_forras or (d.get("csoport") and bool(f["talalat"]))

        if not self.index_kesz:
            self.csoport_kiir("(fájlok beolvasása folyamatban...)")
        elif csoportok:
            self.csoport_kiir(", ".join(csoportok))
        else:
            self.csoport_kiir("nem állapítható meg")

    def csoport_kiir(self, szoveg):
        r = self.kimenet.tag_ranges("csoportblokk")
        if r:
            self.kimenet.delete(r[0], r[-1])
        if not self.aktualis:
            return
        self.kimenet.insert("end", "\n", "csoportblokk")
        self.kimenet.insert("end", f"{CFG.get('csoport_felirat', 'Csoport')} : ", ("csoportblokk", "cim"))
        self.kimenet.insert("end", szoveg, ("csoportblokk", "kiemel"))

    def megnyit(self, f):
        if not f["talalat"]:
            return
        sel = f["lista"].curselection()
        fajl_megnyit(f["talalat"][sel[0] if sel else 0])

    # ---- Gépi programok ----
    def prog_keres(self, e):
        if not self.prog_lista:
            return
        pr, lb = CFG["programok"], self.prog_lista
        lb.delete(0, "end")
        self.prog_talalat = []
        if not self.index_kesz:
            lb.insert("end", "Beolvasás folyamatban...")
            return
        i_kod, i_gep = self.oszlop(pr["kod_fejlec"]), self.oszlop(pr["gep_fejlec"])
        if i_kod is None:
            lb.insert("end", f"Nincs '{pr['kod_fejlec']}' oszlop az Excelben.")
            return
        kodok, bizonytalan = programkodok(e[i_kod])
        if not kodok:
            lb.insert("end", "Nincs programkód megadva.")
            return
        gep = e[i_gep] if i_gep is not None else ""
        jel = "   (? az Excelben)" if bizonytalan else ""
        hianyzo = []
        for eredeti, jeloltek in kodok:
            talalt = [t for j in jeloltek for t in self.prog_idx.get(j, [])]
            sajat = [(n, u) for g, n, u in talalt if g == gep]
            masik = [(n, u) for g, n, u in talalt if g != gep]
            for n, u in sajat:
                lb.insert("end", f"{n}:  {rovid_ut(u)}{jel}")
                self.prog_talalat.append(u)
            if not sajat:
                for n, u in masik:   # nincs a sor gépénél -> más gépnél, jelölve
                    lb.insert("end", f"[másik gép] {n}:  {rovid_ut(u)}{jel}")
                    self.prog_talalat.append(u)
            if not talalt:
                hianyzo.append(eredeti)
        for h in hianyzo:            # az infósorok a megnyitható sorok UTÁN jönnek
            lb.insert("end", f"Nincs program: {h}")
        if self.prog_talalat:
            lb.selection_set(0)

    def prog_megnyit(self):
        sel = self.prog_lista.curselection()
        i = sel[0] if sel else 0
        if i < len(self.prog_talalat):   # "Nincs program" sorra kattintva nem csinál semmit
            fajl_megnyit(self.prog_talalat[i])


if __name__ == "__main__":
    root = tk.Tk()
    if CFG is None:
        root.withdraw()
        messagebox.showerror("Hiba a config.json betöltésénél", f"{APP_MAPPA / 'config.json'}\n\n{_CFG_HIBA}")
    else:
        App(root)
        root.mainloop()
